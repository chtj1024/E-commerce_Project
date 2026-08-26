from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\chlqn\Desktop\shop")
OUT = ROOT / "output" / "resume" / "최태준_백엔드개발자_이력서.docx"
PHOTO = Path(r"C:\Users\chlqn\Desktop\증명사진_최태준.jpg")

# Named design override: KOREAN_TECH_RESUME
# Base: compact_reference_guide. A4 and tighter spacing are deliberate resume overrides.
FONT = "Malgun Gothic"
NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "EAF2F8"
PALE = "F4F6F8"
GRAY = "666666"
DARK = "222222"
WHITE = "FFFFFF"
BORDER = "D9E2F3"
PAGE_WIDTH_DXA = 11906
CONTENT_WIDTH_DXA = 10010


def set_cell_margins(cell, top=70, start=110, bottom=70, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        spec = kwargs[edge]
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        for key in ("val", "sz", "space", "color"):
            if key in spec:
                el.set(qn(f"w:{key}"), str(spec[key]))


def set_table_geometry(table, widths_dxa, indent_dxa=0):
    total = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent_dxa))
    tblInd.set(qn("w:type"), "dxa")
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths_dxa[idx]))
            tcW.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)


def set_run(run, size=9.3, color=DARK, bold=False, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def set_para(p, before=0, after=3, line=1.08, align=None, keep=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep:
        fmt.keep_with_next = True
    return p


def add_hyperlink(paragraph, text, url, color=BLUE, size=8.7, bold=False):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rFonts.set(qn("w:eastAsia"), FONT)
    rPr.append(rFonts)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    rPr.append(color_el)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(size * 2)))
    rPr.append(size_el)
    if bold:
        rPr.append(OxmlElement("w:b"))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    rPr.append(underline)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    set_para(p, before=7, after=4, line=1.0, keep=True)
    set_run(p.add_run(text), size=12.5, color=NAVY, bold=True)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), BLUE)
    borders.append(bottom)
    pPr.append(borders)
    return p


def add_role_header(doc, title, meta, link_text=None, link_url=None):
    p = doc.add_paragraph()
    set_para(p, before=3, after=1, line=1.0, keep=True)
    set_run(p.add_run(title), size=10.3, color=DARK, bold=True)
    set_run(p.add_run(f"  |  {meta}"), size=8.6, color=GRAY)
    if link_text and link_url:
        set_run(p.add_run("  |  "), size=8.6, color=GRAY)
        add_hyperlink(p, link_text, link_url, size=8.6, bold=True)
    return p


def add_bullet(doc, text, num_id, level=0, size=9.0, color=DARK, after=2):
    p = doc.add_paragraph()
    set_para(p, before=0, after=after, line=1.08)
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), str(num_id))
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.insert(0, numPr)
    set_run(p.add_run(text), size=size, color=color)
    return p


def create_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abs_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abs_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    numFmt = OxmlElement("w:numFmt")
    numFmt.set(qn("w:val"), "bullet")
    lvl.append(numFmt)
    lvlText = OxmlElement("w:lvlText")
    lvlText.set(qn("w:val"), "•")
    lvl.append(lvlText)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    pPr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    pPr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "260")
    pPr.append(ind)
    lvl.append(pPr)
    rPr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    rPr.append(fonts)
    lvl.append(rPr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para(p, before=0, after=0, line=1.0)
    set_run(p.add_run("최태준 | Backend Developer    "), size=7.5, color=GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = p.add_run()._r
    run.append(fld_begin)
    run.append(instr)
    run.append(fld_end)
    set_run(p.add_run(" / "), size=7.5, color=GRAY)
    fld_begin2 = OxmlElement("w:fldChar")
    fld_begin2.set(qn("w:fldCharType"), "begin")
    instr2 = OxmlElement("w:instrText")
    instr2.set(qn("xml:space"), "preserve")
    instr2.text = " NUMPAGES "
    fld_end2 = OxmlElement("w:fldChar")
    fld_end2.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()._r
    run2.append(fld_begin2)
    run2.append(instr2)
    run2.append(fld_end2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.2677)
section.page_height = Inches(11.6929)
section.top_margin = Inches(0.52)
section.bottom_margin = Inches(0.50)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)
section.header_distance = Inches(0.24)
section.footer_distance = Inches(0.25)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(9.3)
normal.font.color.rgb = RGBColor.from_string(DARK)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.08

for style_name, size, color, before, after in (
    ("Heading 1", 12.5, NAVY, 7, 4),
    ("Heading 2", 10.3, DARK, 4, 2),
    ("Heading 3", 9.5, BLUE, 3, 1),
):
    st = styles[style_name]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

bullet_num_id = create_bullet_numbering(doc)
add_footer(section)

# First-page header: compact resume variant of customer_pack.
header = doc.add_table(rows=1, cols=2)
set_table_geometry(header, [7840, 2170], indent_dxa=0)
header_cell = header.cell(0, 0)
photo_cell = header.cell(0, 1)
for cell in (header_cell, photo_cell):
    set_cell_margins(cell, top=70, start=100, bottom=70, end=100)
    set_cell_borders(cell,
        top={"val": "nil"}, start={"val": "nil"}, bottom={"val": "nil"}, end={"val": "nil"})
header_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
photo_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

p = header_cell.paragraphs[0]
set_para(p, before=0, after=1, line=1.0)
set_run(p.add_run("최태준"), size=23, color=NAVY, bold=True)
p = header_cell.add_paragraph()
set_para(p, before=0, after=7, line=1.0)
set_run(p.add_run("BACKEND DEVELOPER"), size=10.5, color=BLUE, bold=True)

p = header_cell.add_paragraph()
set_para(p, before=0, after=2, line=1.0)
set_run(p.add_run("010-6866-9066  |  "), size=8.8, color=GRAY)
add_hyperlink(p, "chlqnftkwh@naver.com", "mailto:chlqnftkwh@naver.com", size=8.8)
p = header_cell.add_paragraph()
set_para(p, before=0, after=2, line=1.0)
add_hyperlink(p, "GitHub", "https://github.com/chtj1024", size=8.8, bold=True)
set_run(p.add_run("  |  "), size=8.8, color=GRAY)
add_hyperlink(p, "Portfolio", "https://app.notion.com/p/3c47b079826d81afa1f3f11ad400a673", size=8.8, bold=True)
set_run(p.add_run("  |  "), size=8.8, color=GRAY)
add_hyperlink(p, "Blog", "https://noahchoi.tistory.com/", size=8.8, bold=True)

p = photo_cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para(p, before=0, after=0, line=1.0)
p.add_run().add_picture(str(PHOTO), width=Inches(1.14), height=Inches(1.52))

add_section_heading(doc, "PROFILE")
p = doc.add_paragraph()
set_para(p, before=0, after=4, line=1.12)
set_run(p.add_run(
    "주문·결제 과정의 데이터 정합성을 고민하는 백엔드 개발자입니다. "
    "Spring Boot와 JPA로 이커머스 서비스를 구현하며 조건부 UPDATE와 멱등성 있는 상태 전이로 "
    "동시 주문, 결제 실패 및 주문 만료 상황을 처리했습니다. 핵심 비즈니스 규칙은 실제 MySQL 기반 통합·동시성 테스트로 검증했습니다."
), size=9.3)

add_section_heading(doc, "TECHNICAL SKILLS")
skills = doc.add_table(rows=4, cols=2)
set_table_geometry(skills, [1900, 8110], indent_dxa=0)
skill_rows = [
    ("Backend", "Java 21 · Spring Boot 3.5 · Spring Data JPA · QueryDSL"),
    ("Security / API", "Spring Security · JWT · Swagger / OpenAPI"),
    ("Database / Test", "MySQL 8 · JUnit 5 · Testcontainers"),
    ("Frontend / Tools", "React 19 · TypeScript · Vite · Axios · Git · Gradle · Docker · Codex (AI-assisted development)"),
]
for i, (label, value) in enumerate(skill_rows):
    left, right = skills.rows[i].cells
    for cell in (left, right):
        set_cell_margins(cell, top=58, start=105, bottom=58, end=105)
        set_cell_borders(cell,
            top={"val": "single", "sz": "4", "color": BORDER},
            start={"val": "single", "sz": "4", "color": BORDER},
            bottom={"val": "single", "sz": "4", "color": BORDER},
            end={"val": "single", "sz": "4", "color": BORDER})
    set_cell_shading(left, LIGHT_BLUE)
    lp = left.paragraphs[0]
    set_para(lp, before=0, after=0, line=1.0)
    set_run(lp.add_run(label), size=8.6, color=NAVY, bold=True)
    rp = right.paragraphs[0]
    set_para(rp, before=0, after=0, line=1.0)
    set_run(rp.add_run(value), size=8.7)

add_section_heading(doc, "PROJECT EXPERIENCE")
add_role_header(
    doc,
    "Shop - 주문·결제 데이터 정합성을 고려한 이커머스 서비스",
    "1인 프로젝트 · 2026.04 - 2026.07",
    "GitHub",
    "https://github.com/chtj1024/E-commerce_Project",
)
p = doc.add_paragraph()
set_para(p, before=0, after=3, line=1.08)
set_run(p.add_run(
    "회원·상품·장바구니·주문·결제 흐름을 구현하고, 동시 주문과 결제 상태 경합에서 발생하는 데이터 정합성 문제를 중점적으로 해결했습니다."
), size=8.9, color=GRAY)

add_bullet(doc, "상품 상태와 남은 재고를 조건으로 하는 원자적 UPDATE를 적용해 재고 초과 판매를 방지했습니다.", bullet_num_id)
add_bullet(doc, "복수 상품을 ID 순서로 차감해 자원 점유 순서를 고정하고, 일부 상품 실패 시 전체 재고 차감과 주문 생성을 롤백했습니다.", bullet_num_id)
add_bullet(doc, "PAYMENT_PENDING 상태의 조건부 전이로 결제 성공·실패·만료의 중복 처리와 재고 중복 복구를 방지했습니다.", bullet_num_id)
add_bullet(doc, "30초 주기·최대 100건의 만료 주문 처리와 status + expiresAt 복합 인덱스를 적용하고, 개별 실패를 격리했습니다.", bullet_num_id)
add_bullet(doc, "Testcontainers MySQL 8.4에서 재고 10개 상품에 20개 동시 주문을 실행해 성공 주문 10건과 최종 재고 0개를 검증했습니다.", bullet_num_id)
add_bullet(doc, "QueryDSL 복합 검색과 JWT Access·Refresh Token 분리, Refresh Token Rotation 및 역할 기반 인가를 구현했습니다.", bullet_num_id)

spacer = doc.add_paragraph()
set_para(spacer, before=1, after=1, line=1.0)

add_section_heading(doc, "OTHER PROJECTS")
add_role_header(doc, "출산율 미래 데이터 예측", "1인 프로젝트 · 2022.06 - 2022.12", "GitHub", "https://github.com/chtj1024/BirthRate_Forecast")
add_bullet(doc, "공공데이터를 수집·전처리하고 EDA를 수행한 뒤 ARIMA 시계열 모델로 출산율을 예측하고 결과를 시각화했습니다.", bullet_num_id)

add_role_header(doc, "웹·데스크톱 캘린더", "1인 프로젝트 · 2023.06 - 2023.12", "GitHub", "https://github.com/chtj1024/Calendar_Project")
add_bullet(doc, "React와 Electron으로 일정 CRUD, 드래그 기반 일정 등록, Tray 실행 및 사용자 설정 저장 기능을 구현했습니다.", bullet_num_id)

add_section_heading(doc, "ADDITIONAL EXPERIENCE")
add_role_header(doc, "포트나이트 프로게이머 · 인게임 리더", "루나틱하이 · Cloud9 · 1년 6개월")
p = doc.add_paragraph()
set_para(p, before=0, after=3, line=1.08)
set_run(p.add_run(
    "듀오와 스쿼드 경기에서 인게임 리더를 맡아 실시간 상황 판단, 팀 전략 조율, 경기 분석과 반복 피드백을 수행했습니다."
), size=8.9, color=GRAY)
add_bullet(doc, "중국 DouYu 주최 듀오 대회 2위(상금 800만 원)로 상하이 Stan Lee's Comic Universe Fortnite Open 시드를 획득해 출전했습니다.", bullet_num_id)
add_bullet(doc, "인벤 주최 솔로 대회 80명 중 종합 2위, 듀오 대회 30팀 중 종합 2위를 기록했습니다.", bullet_num_id)
add_bullet(doc, "Cloud9 소속 프로게이머로 부산 G-STAR에 초청되어 인플루언서들과 이벤트 경기에 참여했습니다.", bullet_num_id)

add_section_heading(doc, "EDUCATION")
add_role_header(doc, "인하공업전문대학교 컴퓨터정보과", "인천")
add_bullet(doc, "학사학위 전공심화과정 졸업 · 2024", bullet_num_id)
add_bullet(doc, "전문학사 졸업 · 2023", bullet_num_id)

add_section_heading(doc, "ADDITIONAL INFORMATION")
info = doc.add_table(rows=2, cols=2)
set_table_geometry(info, [5005, 5005], indent_dxa=0)
info_values = [
    ("병역", "군필"),
    ("지원 가능 시점", "즉시"),
]
for i, (label, value) in enumerate(info_values):
    cell = info.cell(i, 0)
    set_cell_margins(cell, top=70, start=110, bottom=70, end=110)
    set_cell_shading(cell, PALE)
    set_cell_borders(cell,
        top={"val": "single", "sz": "4", "color": BORDER},
        start={"val": "single", "sz": "4", "color": BORDER},
        bottom={"val": "single", "sz": "4", "color": BORDER},
        end={"val": "single", "sz": "4", "color": BORDER})
    p = cell.paragraphs[0]
    set_para(p, before=0, after=0, line=1.0)
    set_run(p.add_run(f"{label}  "), size=8.8, color=NAVY, bold=True)
    set_run(p.add_run(value), size=8.8)
    right = info.cell(i, 1)
    set_cell_margins(right, top=70, start=110, bottom=70, end=110)
    set_cell_borders(right,
        top={"val": "nil"}, start={"val": "nil"}, bottom={"val": "nil"}, end={"val": "nil"})
    right.text = ""

doc.core_properties.title = "최태준 백엔드 개발자 이력서"
doc.core_properties.subject = "Java Spring Backend Developer Resume"
doc.core_properties.author = "최태준"
doc.core_properties.keywords = "Java, Spring Boot, Backend, Resume"
doc.core_properties.comments = ""

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
